"""Shared helpers for the small macro-tracking workbook."""

from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
import re
import shutil

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "減脂追蹤.xlsx"
BACKUP_DIR = ROOT / "backups"
TARGET_ROWS = {"Calories": 12, "Protein": 13, "Carbs": 14, "Fat": 15}


def backup_before_edit(workbook_path: Path) -> Path | None:
    """Create a timestamped backup and keep the newest 20 backups."""

    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        return None

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = BACKUP_DIR / f"{workbook_path.stem}_backup_{stamp}.xlsx"
    suffix = 1
    while backup_path.exists():
        backup_path = BACKUP_DIR / f"{workbook_path.stem}_backup_{stamp}_{suffix}.xlsx"
        suffix += 1
    shutil.copy2(workbook_path, backup_path)

    backups = sorted(
        BACKUP_DIR.glob("*.xlsx"),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    for old_backup in backups[20:]:
        old_backup.unlink()
    return backup_path


def load_tracker(workbook_path: Path = DEFAULT_WORKBOOK):
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"找不到追蹤檔案：{workbook_path}")
    return load_workbook(workbook_path)


def save_tracker(workbook, workbook_path: Path = DEFAULT_WORKBOOK) -> None:
    Path(workbook_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)


def set_recalculation(workbook) -> None:
    """Ask Excel to refresh formulas when the workbook is opened."""

    calculation = getattr(workbook, "calculation", None)
    if calculation is not None:
        calculation.calcMode = "auto"
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True


def parse_date(raw: str | None, default: date | None = None) -> date:
    default = default or date.today()
    if not raw:
        return default
    value = raw.strip().lower()
    if value in {"今天", "today"}:
        return date.today()
    if value in {"昨天", "yesterday"}:
        return date.today() - timedelta(days=1)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(raw.strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"日期格式無法辨識：{raw}。請使用 YYYY-MM-DD。")


def date_from_text(text: str, default: date | None = None) -> date:
    match = re.search(r"\b20\d{2}[-/.]\d{1,2}[-/.]\d{1,2}\b", text)
    return parse_date(match.group(0) if match else None, default=default)


def infer_day_type(text: str) -> str | None:
    lowered = text.lower()
    if re.search(r"訓練|訓練日|重訓|運動|workout|training|gym", lowered):
        return "TRAINING"
    if re.search(r"休息|休息日|rest", lowered):
        return "REST"
    return None


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def find_daily_row(worksheet, target_date: date) -> int | None:
    for row in range(5, worksheet.max_row + 1):
        if _as_date(worksheet.cell(row, 1).value) == target_date:
            return row
    return None


def target_min(worksheet, day_type: str, metric: str) -> float:
    row = TARGET_ROWS[metric]
    column = 2 if day_type == "TRAINING" else 4
    return float(worksheet.cell(row, column).value)


def write_daily_formulas(worksheet, row: int) -> None:
    """Write target, difference, and advice formulas for one daily row."""

    refs = {"Calories": "D", "Protein": "G", "Carbs": "J", "Fat": "M"}
    actuals = {"Calories": "C", "Protein": "F", "Carbs": "I", "Fat": "L"}
    diffs = {"Calories": "E", "Protein": "H", "Carbs": "K", "Fat": "N"}

    for metric, target_column in refs.items():
        target_row = TARGET_ROWS[metric]
        worksheet[f"{target_column}{row}"] = (
            f'=IF($B{row}="","",IF($B{row}="TRAINING",'
            f"'設定'!$B${target_row},IF($B{row}=\"REST\",'設定'!$D${target_row},\"\")))"
        )
        actual_column = actuals[metric]
        diff_column = diffs[metric]
        worksheet[f"{diff_column}{row}"] = (
            f'=IF(OR(${actual_column}{row}="",${target_column}{row}=""),"",'
            f"${target_column}{row}-${actual_column}{row})"
        )

    worksheet[f"O{row}"] = (
        f'=IF($A{row}="","",IF($B{row}="","請先填 Day_Type（TRAINING 或 REST）",'
        f'IF(OR($C{row}="",$F{row}="",$I{row}="",$L{row}=""),'
        f'"請補齊 Calories、Protein、Carbs、Fat",'
        f'IF($H{row}>0,"優先補蛋白質。",'
        f'IF($K{row}>0,"蛋白質已足夠，優先補碳水，例如白飯或香蕉。",'
        f'IF($N{row}>0,"優先補脂肪。",'
        f'IF($E{row}>0,"三大營養素已達最低，熱量尚差；可補少量主食。",'
        f'"已達每日最低目標；蛋白質足夠，不需補充乳清。")))))))'
    )


def update_daily_table_ref(worksheet) -> None:
    for table in worksheet.tables.values():
        if table.name == "DailyRecordTable":
            table.ref = f"A4:P{max(4, worksheet.max_row)}"
            return
