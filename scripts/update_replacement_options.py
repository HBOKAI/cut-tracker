"""Add flexible food replacement options to the menu reference sheet."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
import sys

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from add_menu_plan import REPLACEMENT_ROWS, style_headers, style_section, write_rows
from tracker_utils import DEFAULT_WORKBOOK, backup_before_edit, set_recalculation


def main() -> int:
    parser = argparse.ArgumentParser(description="加入 711 飯糰、雞胗、地瓜替換方案")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    if not workbook_path.exists():
        print(f"找不到追蹤檔案：{workbook_path}")
        return 2

    workbook = load_workbook(workbook_path)
    if "菜單規劃" not in workbook.sheetnames:
        print("找不到「菜單規劃」工作表。")
        return 2
    menu = workbook["菜單規劃"]

    if (
        menu["A46"].value == "替換方案"
        and menu["A48"].value == REPLACEMENT_ROWS[0][0]
        and menu["D48"].alignment.wrap_text
        and menu.row_dimensions[48].height == 48
    ):
        print("替換方案已存在，未變更檔案。")
        return 0

    backup_path = backup_before_edit(workbook_path)
    style_section(menu, 46, "H", "替換方案")
    headers = ["替換食物", "可替代", "建議份量", "使用方式", "蛋白質", "碳水", "脂肪", "備註"]
    for column, header in enumerate(headers, 1):
        menu.cell(47, column, header)
    style_headers(menu, 47, 8)
    write_rows(menu, 48, REPLACEMENT_ROWS, 8)
    for row in range(48, 51):
        for column in (4, 8):
            alignment = copy(menu.cell(row, column).alignment)
            alignment.wrap_text = True
            menu.cell(row, column).alignment = alignment
        menu.row_dimensions[row].height = 48

    set_recalculation(workbook)
    workbook.save(workbook_path)
    print(f"已加入替換方案：{workbook_path}")
    if backup_path:
        print(f"已建立備份：{backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
