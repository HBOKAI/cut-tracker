"""Update the rest-day menu to a three-meal schedule."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
import sys

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from tracker_utils import DEFAULT_WORKBOOK, backup_before_edit, set_recalculation


REST_ROWS = [
    ["12:30", "午餐", "150 g（半盒）", "200 g", "約 0.60 杯", "鯖魚日加台糖水煮鯖魚 1/2 罐", "300 g", "鯖魚每週 3 次即可"],
    ["16:00", "下午茶", "—", "100 g", "約 0.30 杯", "全蛋 2 顆＋樂維根麵茶植物蛋白 1–1.5 份", "—", "把原早餐內容移到下午茶；若蛋白質不足再補 0.5 份"],
    ["18:30", "晚餐", "150 g（半盒）", "150 g", "約 0.45 杯", "全蛋 2 顆", "300 g", ""],
]


def write_rows(worksheet, start_row: int) -> None:
    for row_number, values in enumerate(REST_ROWS, start_row):
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(row_number, column, value)
            alignment = copy(cell.alignment)
            alignment.vertical = "center"
            alignment.wrap_text = column in {6, 8}
            cell.alignment = alignment
        worksheet.row_dimensions[row_number].height = 30


def main() -> int:
    parser = argparse.ArgumentParser(description="將休息日菜單改為午餐、下午茶、晚餐")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    if not workbook_path.exists():
        print(f"找不到追蹤檔案：{workbook_path}")
        return 2

    workbook = load_workbook(workbook_path)
    if "菜單規劃" not in workbook.sheetnames:
        print("找不到「菜單規劃」工作表。")
        return 2

    menu = workbook["菜單規劃"]
    already_updated = (
        menu["A17"].value == "12:30"
        and menu["B17"].value == "午餐"
        and menu["A18"].value == "16:00"
        and menu["B18"].value == "下午茶"
        and menu["A19"].value == "18:30"
        and menu["B19"].value == "晚餐"
        and menu["A20"].value == "每日合計"
        and menu["A21"].value in (None, "")
        and menu["A21"].fill.fill_type is None
    )
    if already_updated:
        print("休息日三餐菜單已是最新版本，未變更檔案。")
        return 0

    backup_path = backup_before_edit(workbook_path)

    # Preserve the existing table styling while moving the total row up one row.
    for column in range(1, 9):
        menu.cell(20, column)._style = copy(menu.cell(21, column)._style)
        menu.cell(20, column).number_format = menu.cell(21, column).number_format
    write_rows(menu, 17)

    totals = {
        1: "每日合計",
        3: "300 g（1盒）",
        4: "450 g",
        5: "約 1.35–1.4 杯",
        6: "全蛋 4 顆＋樂維根麵茶植物蛋白 1–1.5 份",
        7: "約 600 g",
    }
    for column in range(1, 9):
        cell = menu.cell(20, column)
        cell.value = totals.get(column, "")
        alignment = copy(cell.alignment)
        alignment.vertical = "center"
        alignment.wrap_text = column == 6
        cell.alignment = alignment
    menu.row_dimensions[20].height = 34

    for row in (21, 22, 23):
        for column in range(1, 9):
            menu.cell(row, column).value = None
            if row == 21:
                menu.cell(row, column)._style = copy(menu.cell(22, column)._style)
    menu.row_dimensions[21].height = None

    set_recalculation(workbook)
    workbook.save(workbook_path)
    print(f"已更新休息日三餐菜單：{workbook_path}")
    if backup_path:
        print(f"已建立備份：{backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
