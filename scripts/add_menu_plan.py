"""Add the user's original weekly meal plan as a reference sheet."""

from __future__ import annotations

import argparse
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from tracker_utils import DEFAULT_WORKBOOK, backup_before_edit, set_recalculation


NAVY = "1F4E79"
BLUE = "5B9BD5"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GREY = "666666"
THIN_GREY = Side(style="thin", color="D9E2F3")

TRAINING_ROWS = [
    ["07:30", "早餐", "—", "120 g", "約 1/3 杯", "全蛋 2 顆", "—", "正常吃，不需刻意低碳"],
    ["12:30", "午餐", "150 g（半盒）", "180 g", "約 0.55 杯", "鯖魚日加台糖水煮鯖魚 1/2 罐", "200 g", "鯖魚每週 3 次即可"],
    ["17:30", "訓前餐", "—", "180 g", "約 0.55 杯", "乳清 1 scoop＋香蕉 1 根", "—", "盡量距訓練 1.5–2.5 小時"],
    ["19:30–21:00", "重訓", "—", "—", "—", "水", "—", "依實際訓練時間前後平移"],
    ["21:00", "訓後晚餐", "150 g（半盒）", "160 g", "約 1/2 杯", "全蛋 2 顆", "200 g", "晚間吃白飯沒有問題"],
    ["睡前", "視需要補蛋白", "—", "—", "—", "乳清 0.5–1 scoop", "—", "當天蛋白質足夠可省略"],
]

REST_ROWS = [
    ["07:30", "早餐", "—", "100 g", "約 0.30 杯", "全蛋 2 顆＋乳清 1 scoop", "—", "休息日仍保留碳水"],
    ["12:30", "午餐", "150 g（半盒）", "200 g", "約 0.60 杯", "鯖魚日加台糖水煮鯖魚 1/2 罐", "300 g", "鯖魚每週 3 次即可"],
    ["18:30", "晚餐", "150 g（半盒）", "150 g", "約 0.45 杯", "全蛋 2 顆", "300 g", ""],
    ["21:30", "睡前", "—", "—", "—", "乳清 0.5–1 scoop", "—", "依當天蛋白質攝取調整"],
]

WEEKLY_ROWS = [
    ["週一", "訓練日", "1盒", "約2杯", "1/2罐", "4顆"],
    ["週二", "訓練日", "1盒", "約2杯", "—", "4顆"],
    ["週三", "休息日", "1盒", "約1.4杯", "—", "4顆"],
    ["週四", "訓練日", "1盒", "約2杯", "1/2罐", "4顆"],
    ["週五", "休息日", "1盒", "約1.4杯", "—", "4顆"],
    ["週六", "訓練日", "1盒", "約2杯", "1/2罐", "4顆"],
    ["週日", "休息日", "1盒", "約1.4杯", "—", "4顆"],
]

SHOPPING_ROWS = [
    ["雞胸 300 g/盒", "7 盒"],
    ["台糖水煮鯖魚", "2 罐即可（實吃 1.5 罐）"],
    ["雞蛋", "28 顆"],
    ["生米", "約 12.2 米杯/週"],
    ["香蕉", "至少 4 根"],
    ["青菜", "約 3.4 kg/週以上"],
    ["乳清", "約 11–14 scoop/週"],
]


def style_title(worksheet, end_column: str, title: str, note: str) -> None:
    worksheet.merge_cells(f"A1:{end_column}1")
    worksheet["A1"] = title
    worksheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 28

    worksheet.merge_cells(f"A2:{end_column}2")
    worksheet["A2"] = note
    worksheet["A2"].font = Font(name="Aptos", size=10, italic=True, color=GREY)
    worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[2].height = 30


def style_section(worksheet, row: int, end_column: str, text: str) -> None:
    worksheet.merge_cells(f"A{row}:{end_column}{row}")
    cell = worksheet[f"A{row}"]
    cell.value = text
    cell.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[row].height = 22


def style_headers(worksheet, row: int, count: int) -> None:
    for col in range(1, count + 1):
        cell = worksheet.cell(row, col)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=WHITE))
    worksheet.row_dimensions[row].height = 28


def write_rows(worksheet, start_row: int, rows: list[list[str]], column_count: int) -> int:
    for row_number, values in enumerate(rows, start_row):
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(row_number, column, value)
            cell.border = Border(bottom=THIN_GREY)
            cell.alignment = Alignment(vertical="center", wrap_text=column in {6, 8})
        worksheet.row_dimensions[row_number].height = 30 if column_count == 8 else 22
    return start_row + len(rows) - 1


def build_menu_sheet(workbook):
    if "菜單規劃" in workbook.sheetnames:
        return None

    worksheet = workbook.create_sheet("菜單規劃")
    worksheet.sheet_view.showGridLines = False
    style_title(
        worksheet,
        "H",
        "減脂菜單規劃｜一週四練",
        "保留原始菜單作為參考；日常只需在「每日紀錄」填入 Calories、Protein、Carbs、Fat。",
    )

    menu_headers = ["時間", "餐次", "雞胸（生重）", "白飯（熟重）", "生米約米杯", "其他食物", "青菜", "備註"]
    style_section(worksheet, 4, "H", "訓練日菜單｜晚上重訓")
    for col, header in enumerate(menu_headers, 1):
        worksheet.cell(5, col, header)
    style_headers(worksheet, 5, 8)
    training_end = write_rows(worksheet, 6, TRAINING_ROWS, 8)
    total_row = training_end + 1
    worksheet.cell(total_row, 1, "每日合計")
    worksheet.cell(total_row, 3, "300 g（1盒）")
    worksheet.cell(total_row, 4, "640 g")
    worksheet.cell(total_row, 5, "約 1.9–2.0 杯")
    worksheet.cell(total_row, 6, "全蛋 4 顆＋乳清 1.5–2 scoop＋香蕉 1 根")
    worksheet.cell(total_row, 7, "至少 400 g")
    for col in range(1, 9):
        worksheet.cell(total_row, col).font = Font(name="Aptos", size=10, bold=True)
        worksheet.cell(total_row, col).fill = PatternFill("solid", fgColor=YELLOW)
        worksheet.cell(total_row, col).border = Border(bottom=THIN_GREY)
        worksheet.cell(total_row, col).alignment = Alignment(vertical="center", wrap_text=col == 6)
    worksheet.row_dimensions[total_row].height = 34

    rest_section = total_row + 3
    style_section(worksheet, rest_section, "H", "休息日菜單")
    rest_header = rest_section + 1
    for col, header in enumerate(menu_headers, 1):
        worksheet.cell(rest_header, col, header)
    style_headers(worksheet, rest_header, 8)
    rest_start = rest_header + 1
    rest_end = write_rows(worksheet, rest_start, REST_ROWS, 8)
    rest_total = rest_end + 1
    worksheet.cell(rest_total, 1, "每日合計")
    worksheet.cell(rest_total, 3, "300 g（1盒）")
    worksheet.cell(rest_total, 4, "450 g")
    worksheet.cell(rest_total, 5, "約 1.35–1.4 杯")
    worksheet.cell(rest_total, 6, "全蛋 4 顆＋乳清 1.5–2 scoop")
    worksheet.cell(rest_total, 7, "約 600 g")
    for col in range(1, 9):
        worksheet.cell(rest_total, col).font = Font(name="Aptos", size=10, bold=True)
        worksheet.cell(rest_total, col).fill = PatternFill("solid", fgColor=YELLOW)
        worksheet.cell(rest_total, col).border = Border(bottom=THIN_GREY)
        worksheet.cell(rest_total, col).alignment = Alignment(vertical="center", wrap_text=col == 6)
    worksheet.row_dimensions[rest_total].height = 34

    week_section = rest_total + 3
    style_section(worksheet, week_section, "F", "每週安排")
    week_header = week_section + 1
    for col, header in enumerate(["星期", "類型", "雞胸", "生米", "鯖魚", "蛋"], 1):
        worksheet.cell(week_header, col, header)
    style_headers(worksheet, week_header, 6)
    week_end = write_rows(worksheet, week_header + 1, WEEKLY_ROWS, 6)

    shopping_section = week_end + 3
    style_section(worksheet, shopping_section, "B", "每週採買建議")
    shopping_header = shopping_section + 1
    worksheet.cell(shopping_header, 1, "項目")
    worksheet.cell(shopping_header, 2, "數量")
    style_headers(worksheet, shopping_header, 2)
    write_rows(worksheet, shopping_header + 1, SHOPPING_ROWS, 2)

    widths = {"A": 15, "B": 16, "C": 18, "D": 16, "E": 16, "F": 38, "G": 14, "H": 30}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A5"
    return worksheet


def main() -> int:
    parser = argparse.ArgumentParser(description="將原始一週四練菜單加入追蹤檔")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    if not workbook_path.exists():
        print(f"找不到追蹤檔案：{workbook_path}")
        return 2
    workbook = load_workbook(workbook_path)
    if "菜單規劃" in workbook.sheetnames:
        print("菜單規劃頁已存在，未變更檔案。")
        return 0
    backup_path = backup_before_edit(workbook_path)
    build_menu_sheet(workbook)
    set_recalculation(workbook)
    workbook.save(workbook_path)
    print(f"已加入菜單規劃頁：{workbook_path}")
    if backup_path:
        print(f"已建立備份：{backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
