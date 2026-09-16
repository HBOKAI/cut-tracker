"""Use the user's actual protein product name in the menu reference sheet."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
import sys

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from tracker_utils import DEFAULT_WORKBOOK, backup_before_edit, set_recalculation


REPLACEMENTS = {
    "F8": "樂維根麵茶植物蛋白 1 份（40 g）＋香蕉 1 根",
    "F11": "樂維根麵茶植物蛋白 0.5–1 份",
    "F12": "全蛋 4 顆＋樂維根麵茶植物蛋白 1.5–2 份＋香蕉 1 根",
    "F18": "全蛋 2 顆＋樂維根麵茶植物蛋白 1–1.5 份",
    "H18": "把原早餐內容移到下午茶；若蛋白質不足再補 0.5 份",
    "F20": "全蛋 4 顆＋樂維根麵茶植物蛋白 1–1.5 份",
    "D48": "可放午餐或下午茶；蛋白質不足時再補樂維根麵茶植物蛋白或無糖豆漿",
    "A43": "樂維根麵茶植物蛋白",
    "B43": "約 11–14 份/週",
}


def main() -> int:
    parser = argparse.ArgumentParser(description="更新菜單中的實際蛋白產品名稱")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    if not workbook_path.exists():
        print(f"找不到追蹤檔案：{workbook_path}")
        return 2
    workbook = load_workbook(workbook_path)
    if "菜單規劃" not in workbook.sheetnames:
        print("找不到「菜單規劃」工作表。")
        return 2
    menu = workbook["菜單規劃"]
    if all(menu[cell].value == value for cell, value in REPLACEMENTS.items()):
        print("蛋白產品名稱已是最新版本，未變更檔案。")
        return 0

    backup_path = backup_before_edit(workbook_path)
    for cell, value in REPLACEMENTS.items():
        menu[cell] = value
    for cell in ("F8", "F11", "F12", "F18", "F20", "D48", "H18"):
        alignment = copy(menu[cell].alignment)
        alignment.wrap_text = True
        menu[cell].alignment = alignment
    menu.row_dimensions[18].height = 30
    menu.row_dimensions[48].height = 48
    set_recalculation(workbook)
    workbook.save(workbook_path)
    print(f"已更新蛋白產品名稱：{workbook_path}")
    if backup_path:
        print(f"已建立備份：{backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
