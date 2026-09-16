"""Create the intentionally small two-sheet macro tracker."""

from __future__ import annotations

import argparse
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from tracker_utils import DEFAULT_WORKBOOK, backup_before_edit, set_recalculation


NAVY = "1F4E79"
BLUE = "5B9BD5"
YELLOW = "FFF2CC"
AMBER = "FCE4D6"
WHITE = "FFFFFF"
GREY = "666666"
THIN_GREY = Side(style="thin", color="D9E2F3")


def style_table(table: Table) -> None:
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )


def style_title(worksheet, end_column: str, title: str, note: str) -> None:
    worksheet.merge_cells(f"A1:{end_column}1")
    worksheet["A1"] = title
    worksheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 28

    worksheet.merge_cells(f"A2:{end_column}2")
    worksheet["A2"] = note
    worksheet["A2"].font = Font(name="Aptos", size=10, italic=True, color=GREY)
    worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[2].height = 30


def style_section(worksheet, row: int, end_column: str, text: str) -> None:
    worksheet.merge_cells(f"A{row}:{end_column}{row}")
    cell = worksheet[f"A{row}"]
    cell.value = text
    cell.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[row].height = 22


def style_header_row(worksheet, row: int, end_column: int) -> None:
    for col in range(1, end_column + 1):
        cell = worksheet.cell(row, col)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=WHITE))
    worksheet.row_dimensions[row].height = 28


def build_daily_sheet(workbook: Workbook):
    worksheet = workbook.active
    worksheet.title = "每日紀錄"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A5"
    style_title(
        worksheet,
        "P",
        "減脂飲食追蹤｜每日紀錄",
        "每天一列。差額 = 目標最低值 − 實際；正數表示還差多少。可直接用自然語言更新。",
    )

    headers = [
        "日期", "Day_Type", "Calories", "Target_Calories", "Calorie_Diff",
        "Protein", "Target_Protein", "Protein_Diff", "Carbs", "Target_Carbs",
        "Carbs_Diff", "Fat", "Target_Fat", "Fat_Diff", "建議", "備註",
    ]
    for col, header in enumerate(headers, 1):
        worksheet.cell(4, col).value = header
    style_header_row(worksheet, 4, len(headers))

    table = Table(displayName="DailyRecordTable", ref="A4:P4")
    style_table(table)
    worksheet.add_table(table)

    for col in range(1, 17):
        worksheet.cell(5, col).border = Border(bottom=THIN_GREY)
    for cell in ("A5", "B5", "C5", "F5", "I5", "L5", "P5"):
        worksheet[cell].fill = PatternFill("solid", fgColor=YELLOW)

    validation = DataValidation(type="list", formula1='"TRAINING,REST"', allow_blank=True)
    validation.error = "請填 TRAINING 或 REST"
    validation.errorTitle = "Day_Type 不正確"
    validation.prompt = "訓練日填 TRAINING，休息日填 REST"
    validation.promptTitle = "Day_Type"
    worksheet.add_data_validation(validation)
    validation.add("B5:B5000")

    worksheet["A5"].number_format = "yyyy-mm-dd"
    for col in (3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14):
        worksheet.cell(5, col).number_format = "0.0"
    worksheet["O5"].alignment = Alignment(wrap_text=True, vertical="top")
    worksheet["P5"].alignment = Alignment(wrap_text=True, vertical="top")

    for diff_range in ("E5:E5000", "H5:H5000", "K5:K5000", "N5:N5000"):
        worksheet.conditional_formatting.add(
            diff_range,
            CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=AMBER)),
        )

    widths = {
        "A": 13, "B": 12, "C": 12, "D": 15, "E": 13, "F": 11, "G": 14,
        "H": 13, "I": 11, "J": 13, "K": 12, "L": 10, "M": 12, "N": 11,
        "O": 48, "P": 28,
    }
    for col, width in widths.items():
        worksheet.column_dimensions[col].width = width
    worksheet.auto_filter.ref = "A4:P4"
    return worksheet


def build_settings_sheet(workbook: Workbook):
    worksheet = workbook.create_sheet("設定")
    worksheet.sheet_view.showGridLines = False
    style_title(
        worksheet,
        "F",
        "減脂飲食追蹤｜設定",
        "只需調整黃色欄位。每日紀錄會依 Day_Type 自動比較目標最低值。",
    )
    style_section(worksheet, 4, "F", "基本資料（保留目前設定）")
    basic_headers = ["項目", "數值", "單位", "說明"]
    for col, header in enumerate(basic_headers, 1):
        worksheet.cell(5, col).value = header
    style_header_row(worksheet, 5, 4)
    basics = [
        ("身高", 174, "cm", "使用者提供"),
        ("目前體重", 76.7, "kg", "起始資料；可手動更新"),
        ("目前體脂率", 0.235, "%", "起始資料；可手動更新"),
    ]
    for row, values in enumerate(basics, 6):
        for col, value in enumerate(values, 1):
            worksheet.cell(row, col).value = value
            worksheet.cell(row, col).border = Border(bottom=THIN_GREY)
        worksheet[f"B{row}"].fill = PatternFill("solid", fgColor=YELLOW)
    worksheet["B8"].number_format = "0.0%"

    style_section(worksheet, 10, "F", "每日目標")
    target_headers = ["Metric", "TRAINING 下限", "TRAINING 上限", "REST 下限", "REST 上限", "Unit"]
    for col, header in enumerate(target_headers, 1):
        worksheet.cell(11, col).value = header
    style_header_row(worksheet, 11, 6)
    targets = [
        ("Calories", 2150, 2150, 1850, 1850, "kcal"),
        ("Protein", 165, 175, 160, 170, "g"),
        ("Carbs", 220, 240, 150, 165, "g"),
        ("Fat", 55, 65, 55, 65, "g"),
    ]
    for row, values in enumerate(targets, 12):
        for col, value in enumerate(values, 1):
            worksheet.cell(row, col).value = value
            worksheet.cell(row, col).border = Border(bottom=THIN_GREY)
            if col in (2, 3, 4, 5):
                worksheet.cell(row, col).fill = PatternFill("solid", fgColor=YELLOW)
                worksheet.cell(row, col).number_format = "0.0"
    table = Table(displayName="TargetTable", ref="A11:F15")
    style_table(table)
    worksheet.add_table(table)

    worksheet.merge_cells("A17:F17")
    worksheet["A17"] = "差額以每日目標的下限計算；若 Protein 已達下限，建議欄不會推薦乳清。"
    worksheet["A17"].font = Font(name="Aptos", size=10, italic=True, color=GREY)
    worksheet["A17"].alignment = Alignment(wrap_text=True)
    worksheet.row_dimensions[17].height = 24

    for col, width in {"A": 20, "B": 16, "C": 16, "D": 14, "E": 14, "F": 12}.items():
        worksheet.column_dimensions[col].width = width
    worksheet.freeze_panes = "A11"
    return worksheet


def main() -> int:
    parser = argparse.ArgumentParser(description="建立簡化版減脂飲食追蹤檔")
    parser.add_argument("--output", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--force", action="store_true", help="覆寫既有檔案；覆寫前會先備份")
    args = parser.parse_args()

    output = args.output.resolve()
    if output.exists() and not args.force:
        print(f"檔案已存在，未覆寫：{output}")
        print("若要重新建立，請加上 --force；覆寫前會自動備份。")
        return 1
    if output.exists():
        backup_path = backup_before_edit(output)
        if backup_path:
            print(f"已建立備份：{backup_path}")

    workbook = Workbook()
    build_daily_sheet(workbook)
    build_settings_sheet(workbook)
    set_recalculation(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"已建立簡化版追蹤檔：{output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
